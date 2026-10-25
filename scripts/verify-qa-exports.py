"""Independently reopen XLSX/CSV samples and reconcile values and persisted layout."""
import csv
import json
import pathlib
import sys

import openpyxl

output = pathlib.Path(sys.argv[1] if len(sys.argv) > 1 else "artifacts/qa-exports")
results = []
for size in [0, 1, 45, 10000]:
    expected = json.loads((output / f"expected-{size}.json").read_text(encoding='utf-8'))
    book = openpyxl.load_workbook(output / f"attendance-{size}.xlsx", data_only=False)
    assert book.sheetnames == ["Summary", "Attendance", "Audit trail", "Report info"]
    for name, header_key, rows_key, columns, pane in [
        ("Attendance", "attendanceHeaders", "attendanceRows", "I", "C2"),
        ("Audit trail", "auditHeaders", "auditRows", "L", "A2"),
    ]:
        sheet = book[name]
        expected_rows = [expected[header_key], *expected[rows_key]]
        actual_rows = [
            ["" if cell.value is None else cell.value for cell in row]
            for row in sheet.iter_rows()
        ]
        assert actual_rows == expected_rows, f"{size}: {name} values or membership differ"
        assert sheet.freeze_panes == pane, f"{size}: {name} freeze panes missing"
        assert sheet.auto_filter.ref == f"A1:{columns}{len(expected_rows)}"
        assert not sheet.merged_cells.ranges, "Merged cells must not obstruct data-table sorting"
        assert sheet["A1"].font.bold and sheet["A1"].fill.fgColor.rgb == "FF123B32"
        assert sheet.row_dimensions[1].height >= 32
        for row in sheet.iter_rows(min_row=2):
            assert sheet.row_dimensions[row[0].row].height >= 22
            for cell in row:
                assert cell.data_type != "f", f"Formula introduced at {name}!{cell.coordinate}"
                if cell.value is not None:
                    assert cell.data_type == "s", f"Text type lost at {name}!{cell.coordinate}"
                    assert cell.number_format == "@"
                    assert cell.alignment.wrap_text is True
        assert all(d.width > 0 for d in sheet.column_dimensions.values())
    summary = book["Summary"]
    assert summary["B10"].value == size and summary["B10"].data_type == "n"
    assert summary["B16"].value == expected["rate"]
    assert summary["B16"].number_format == "0.0%"
    for row, status in enumerate(["PRESENT", "LATE", "ABSENT", "EXCUSED", "PENDING"], 11):
        assert summary.cell(row, 2).value == expected["counts"][status]
        assert summary.cell(row, 2).data_type == "n"
        assert summary.cell(row, 3).value == (expected["counts"][status] / size if size else 0)
        assert summary.cell(row, 3).number_format == "0.0%"
    with (output / f"attendance-{size}.csv").open(encoding="utf-8-sig", newline="") as source:
        assert list(csv.reader(source)) == [expected["csvHeaders"], *expected["csvRows"]]
    results.append({
        "records": size,
        "auditEvents": len(expected["auditRows"]),
        "allValuesAndMembershipReconciled": True,
        "textTypesAndFormulaSafety": True,
        "freezeWrapHeightsFiltersAndFormatsPersisted": True,
        "reader": f"openpyxl {openpyxl.__version__}",
        "desktopVisualVerification": "Not tested",
    })
report = {"syntheticData": True, "results": results}
(output / "independent-validation.json").write_text(json.dumps(report, indent=2) + "\n")
print(json.dumps(report, indent=2))
